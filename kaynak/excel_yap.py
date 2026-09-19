# -*- coding: utf-8 -*-
"""Mahal listesi Excel'ini veri.py + imalat.py'den üretir."""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from veri import KATLAR, PARSEL, BINA, KAT_YUKSEKLIK, NET_TAVAN, net_alan, kat_ozeti
from imalat import POZ, GENEL, SUTUNLAR, mahal_satirlari

LACI, MAVI, ACIK, GRI, SARI = "1F3864", "2E75B6", "D9E2F3", "F2F2F2", "FFF2CC"
F = "Arial"
thin = Side(style="thin", color="BFBFBF")
kalin = Side(style="medium", color="1F3864")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def baslik(ws, metin, satir, son, boy=14):
    ws.merge_cells(start_row=satir, start_column=1, end_row=satir, end_column=son)
    c = ws.cell(row=satir, column=1, value=metin)
    c.font = Font(name=F, size=boy, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=LACI)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[satir].height = boy * 2.1

def altbaslik(ws, metin, satir, son):
    ws.merge_cells(start_row=satir, start_column=1, end_row=satir, end_column=son)
    c = ws.cell(row=satir, column=1, value=metin)
    c.font = Font(name=F, size=11, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=MAVI)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[satir].height = 20

def tbaslik(ws, basliklar, satir):
    for i, b in enumerate(basliklar, 1):
        c = ws.cell(row=satir, column=i, value=b)
        c.font = Font(name=F, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=MAVI)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(left=thin, right=thin, top=kalin, bottom=kalin)
    ws.row_dimensions[satir].height = 30

def satir_yaz(ws, s, degerler, boy=34, punto=9, orta=(), kalin_sut=()):
    for i, v in enumerate(degerler, 1):
        c = ws.cell(row=s, column=i, value=v)
        c.font = Font(name=F, size=punto, bold=(i in kalin_sut))
        c.border = BOX
        c.alignment = Alignment(vertical="center", wrap_text=True,
                                horizontal="center" if i in orta else "left",
                                indent=0 if i in orta else 1)
    ws.row_dimensions[s].height = boy

def yaz(hedef):
    wb = openpyxl.Workbook()

    # ---------------------------------------------------------- Genel Bilgiler
    ws = wb.active; ws.title = "Genel Bilgiler"; ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 36; ws.column_dimensions["B"].width = 62
    ws.column_dimensions["C"].width = 46
    baslik(ws, "MAHAL LİSTESİ, İMALAT TARİFİ VE MARKA SEÇİMLERİ", 1, 3, 15)
    ws.merge_cells("A2:C2")
    c = ws["A2"]; c.value = "Durunday Villa — Konya / Meram / Durunday · Bodrumlu dubleks villa"
    c.font = Font(name=F, size=12, bold=True, color=LACI); c.alignment = Alignment(horizontal="center")
    s = 4
    altbaslik(ws, "PROJE VE İMAR KÜNYESİ", s, 3); s += 1
    ozet = kat_ozeti()
    kapali = sum(v[0] for v in ozet.values()); acik = sum(v[1] for v in ozet.values())
    for k, v in [
        ("Yeri", "Konya, Meram ilçesi, Durunday mahallesi"),
        ("Yapı türü", "Ayrık nizam, tek aileli müstakil villa (dubleks)"),
        ("Parsel alanı", "%.0f m² (%.2f × %.2f m — varsayım)" % (PARSEL["alan"], PARSEL["en"], PARSEL["boy"])),
        ("TAKS / KAKS", "%.2f / %.2f  →  taban %.0f m², emsal %.0f m²" % (
            PARSEL["taks"], PARSEL["kaks"], PARSEL["alan"]*PARSEL["taks"], PARSEL["alan"]*PARSEL["kaks"])),
        ("Çekme mesafeleri", "ön %.1f m · yan %.1f m · arka %.1f m" % (
            PARSEL["cekme_on"], PARSEL["cekme_yan"], PARSEL["cekme_arka"])),
        ("Bina oturumu", "%.2f × %.2f m = %.0f m²" % (BINA["en"], BINA["boy"], BINA["en"]*BINA["boy"])),
        ("Kat adedi", "Bodrum + Zemin + 1. Normal kat (dubleks) + çatı arası"),
        ("Kat yükseklikleri", " · ".join("%s %.2f m (net %.2f m)" % (k2, KAT_YUKSEKLIK[k2], NET_TAVAN[k2]) for k2 in KAT_YUKSEKLIK)),
        ("Toplam brüt inşaat alanı", "900 m² (3 kat × 300 m²)"),
        ("Toplam net alan", "%.1f m² kapalı + %.1f m² açık teras/balkon" % (kapali, acik)),
        ("Kalite segmenti", "Üst segment / lüks konut"),
        ("Belge tarihi / revizyon", "19.09.2026 · R01 (çizimlerle birlikte)"),
    ]:
        ws.cell(row=s, column=1, value=k).font = Font(name=F, size=10, bold=True)
        ws.cell(row=s, column=1).fill = PatternFill("solid", fgColor=GRI)
        ws.merge_cells(start_row=s, start_column=2, end_row=s, end_column=3)
        ws.cell(row=s, column=2, value=v).font = Font(name=F, size=10)
        for col in (1, 2, 3):
            ws.cell(row=s, column=col).border = BOX
            ws.cell(row=s, column=col).alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        ws.row_dimensions[s].height = 22
        s += 1
    s += 1
    altbaslik(ws, "GENEL YAPIM TANIMLARI VE MARKA SEÇİMLERİ", s, 3); s += 1
    tbaslik(ws, ["İmalat", "Teknik tanım", "Marka / ürün seçimi"], s); s += 1
    for ad, tanim, marka in GENEL:
        satir_yaz(ws, s, [ad, tanim, marka], boy=46, punto=10, kalin_sut=(1,))
        s += 1
    s += 1
    altbaslik(ws, "VARSAYIMLAR VE UYARILAR", s, 3); s += 1
    for n in [
        "Çizimler ön tasarım (avan) niteliğindedir; uygulama projesi, statik-mekanik-elektrik projeleri ve ruhsat için yetkili proje müellifi onayı gerekir.",
        "Parsel ölçüleri, TAKS/KAKS, çekme mesafesi ve kat adedi Meram Belediyesi imar durumu belgesiyle teyit edilmelidir.",
        "Mahal alanları duvar aksı geometrisinden net (duvar içi) olarak hesaplanmıştır; uygulamada ±%2 sapma olabilir.",
        "Piyes ölçüleri Planlı Alanlar İmar Yönetmeliği Madde 29 asgari değerleriyle karşılaştırılmıştır (bkz. Yönetmelik Uygunluk sayfası).",
        "Konya TS 825'e göre 3. iklim bölgesindedir; yalıtım kalınlıkları buna göre verilmiştir. Deprem hesabı TBDY-2018'e göre yapılacaktır.",
        "Marka seçimleri üst segment için örnek niteliğindedir; teklif aşamasında eşdeğer ürünler önerilebilir.",
    ]:
        ws.merge_cells(start_row=s, start_column=1, end_row=s, end_column=3)
        c = ws.cell(row=s, column=1, value="•  " + n)
        c.font = Font(name=F, size=10); c.fill = PatternFill("solid", fgColor=SARI)
        c.alignment = Alignment(wrap_text=True, vertical="center", indent=1); c.border = BOX
        ws.row_dimensions[s].height = 30
        s += 1
    ws.page_setup.paperSize = 9; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0

    # ---------------------------------------------------------- Mahal Programı
    ws = wb.create_sheet("Mahal Programı"); ws.sheet_view.showGridLines = False
    for i, w in enumerate([11, 44, 14, 12, 12, 16, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    baslik(ws, "MAHAL PROGRAMI — ALAN TABLOSU", 1, 7)
    s = 3
    tbaslik(ws, ["Mahal No", "Mahal Adı", "Net Alan (m²)", "Dar Kenar (m)", "Uzun Kenar (m)", "Tip", "Açıklama"], s)
    s += 1; ws.freeze_panes = "A4"
    toplam_satir = []
    for kat, mahaller in KATLAR.items():
        altbaslik(ws, "%s KAT" % kat.upper() if kat != "1. Kat" else "1. NORMAL KAT", s, 7); s += 1
        ilk = s
        for r in mahaller:
            a, dar, uzun = net_alan(r)
            satir_yaz(ws, s, [r[0], r[1], a, dar, uzun, r[6],
                              "Açık alan — kapalı alana dahil değil" if r[6] == "acik" else ""],
                      boy=18, punto=10, orta=(1, 3, 4, 5, 6), kalin_sut=(1, 2))
            ws.cell(row=s, column=3).number_format = "#,##0.0"
            s += 1
        ws.cell(row=s, column=2, value="%s — TOPLAM" % kat)
        ws.cell(row=s, column=3, value="=SUM(C%d:C%d)" % (ilk, s-1))
        for col in range(1, 8):
            c = ws.cell(row=s, column=col)
            c.font = Font(name=F, size=10, bold=True, color=LACI)
            c.fill = PatternFill("solid", fgColor=ACIK); c.border = BOX
            c.alignment = Alignment(vertical="center", horizontal="center" if col in (1,3,4,5,6) else "left", indent=1)
        ws.cell(row=s, column=3).number_format = "#,##0.0"
        toplam_satir.append(s); s += 2
    altbaslik(ws, "GENEL TOPLAM", s, 7); s += 1
    for ad, deg in [("Toplam net alan (kapalı + açık)", "=" + "+".join("C%d" % r for r in toplam_satir)),
                    ("Toplam brüt inşaat alanı", 900.0),
                    ("Oturum (taban) alanı", float(BINA["en"] * BINA["boy"]))]:
        ws.cell(row=s, column=2, value=ad); ws.cell(row=s, column=3, value=deg)
        for col in range(1, 8):
            c = ws.cell(row=s, column=col)
            c.font = Font(name=F, size=11, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=LACI); c.border = BOX
            c.alignment = Alignment(vertical="center", horizontal="center" if col in (1,3,4,5,6) else "left", indent=1)
        ws.cell(row=s, column=3).number_format = "#,##0.0"
        s += 1
    ws.page_setup.paperSize = 9; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.print_title_rows = "3:3"

    # ---------------------------------------------------------- Mahal Listesi
    ws = wb.create_sheet("Mahal Listesi"); ws.sheet_view.showGridLines = False
    gen = [9, 30, 9] + [13, 11, 21, 12, 17, 16, 13, 17, 13] + [44]
    for i, w in enumerate(gen, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    N = len(gen)
    baslik(ws, "MAHAL LİSTESİ — DURUNDAY VİLLA", 1, N)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=N)
    c = ws.cell(row=2, column=1, value="Poz kodlarının tanımı ve marka seçimleri için «İmalat Tanımları» sayfasına bakınız.")
    c.font = Font(name=F, size=9, italic=True, color="404040"); c.alignment = Alignment(horizontal="center")
    s = 4
    tbaslik(ws, ["Mahal No", "Mahal Adı", "Net Alan (m²)"] + SUTUNLAR + ["Açıklama / Özel Notlar"], s)
    s += 1; ws.freeze_panes = "C5"
    satirlar = mahal_satirlari()
    alanlar = []
    son_kat = None
    for kat, kod, ad, tip, alan, dar, uzun, degerler, aciklama in satirlar:
        if kat != son_kat:
            altbaslik(ws, "%s KAT" % kat.upper() if kat != "1. Kat" else "1. NORMAL KAT", s, N); s += 1
            son_kat = kat
        satir_yaz(ws, s, [kod, ad, alan] + degerler + [aciklama], boy=32, punto=9, orta=(1, 3), kalin_sut=(1, 2))
        ws.cell(row=s, column=3).number_format = "#,##0.0"
        ws.cell(row=s, column=1).font = Font(name=F, size=9, bold=True, color=LACI)
        ws.cell(row=s, column=N).font = Font(name=F, size=8, color="404040")
        alanlar.append(s); s += 1
    ws.cell(row=s, column=2, value="TOPLAM NET ALAN")
    ws.cell(row=s, column=3, value="=SUM(%s)" % ",".join("C%d" % r for r in alanlar))
    for col in range(1, N+1):
        c = ws.cell(row=s, column=col)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=LACI); c.border = BOX
        c.alignment = Alignment(vertical="center", horizontal="center" if col in (1,3) else "left", indent=1)
    ws.cell(row=s, column=3).number_format = "#,##0.0"
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = 8
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "4:4"; ws.page_margins = PageMargins(left=.3, right=.3, top=.4, bottom=.4)

    # ---------------------------------------------------------- İmalat Tanımları
    ws = wb.create_sheet("İmalat Tanımları"); ws.sheet_view.showGridLines = False
    for i, w in enumerate([11, 32, 74, 58], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    baslik(ws, "İMALAT TANIMLARI, ŞARTNAME VE MARKA SEÇİMLERİ", 1, 4)
    s = 3
    tbaslik(ws, ["Kod", "İmalat", "Teknik tanım / şartname", "Marka / ürün seçimi (üst segment)"], s)
    s += 1; ws.freeze_panes = "A4"
    gruplar = [("DÖŞEME KAPLAMASI", "DK"), ("SÜPÜRGELİK", "SP"), ("DUVAR", "DV"), ("TAVAN", "TV"),
               ("KAPI", "KP"), ("DOĞRAMA / PENCERE", "DG"), ("ELEKTRİK", "EL"),
               ("MEKANİK", "MK"), ("SIHHİ TESİSAT", "ST")]
    for grup_ad, on in gruplar:
        altbaslik(ws, grup_ad, s, 4); s += 1
        for kod in sorted(k for k in POZ if k.startswith(on)):
            ad, tanim, marka = POZ[kod]
            satir_yaz(ws, s, [kod, ad, tanim, marka], boy=42, punto=10, orta=(1,), kalin_sut=(1, 2))
            ws.cell(row=s, column=1).font = Font(name=F, size=10, bold=True, color=LACI)
            ws.cell(row=s, column=4).font = Font(name=F, size=10, color="1F3864")
            s += 1
        s += 1
    ws.page_setup.paperSize = 9; ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.print_title_rows = "3:3"

    # ---------------------------------------------------------- Yönetmelik
    ws = wb.create_sheet("Yönetmelik Uygunluk"); ws.sheet_view.showGridLines = False
    for i, w in enumerate([40, 26, 26, 14, 44], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    baslik(ws, "YÖNETMELİK UYGUNLUK KONTROLÜ", 1, 5)
    s = 3
    tbaslik(ws, ["Konu", "Yönetmelik asgarisi", "Projede", "Durum", "Dayanak / açıklama"], s)
    s += 1
    from yonetmelik import KONTROLLER
    for konu, asgari, projede, durum, dayanak in KONTROLLER():
        satir_yaz(ws, s, [konu, asgari, projede, durum, dayanak], boy=30, punto=10, orta=(4,), kalin_sut=(1,))
        c = ws.cell(row=s, column=4)
        c.fill = PatternFill("solid", fgColor="C6EFCE" if durum == "UYGUN" else "FFEB9C")
        c.font = Font(name=F, size=10, bold=True, color="006100" if durum == "UYGUN" else "9C6500")
        s += 1
    ws.page_setup.paperSize = 9; ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.print_title_rows = "3:3"

    wb.save(hedef)
    return hedef

if __name__ == "__main__":
    kok = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    p = yaz(os.path.join(kok, "Durunday Villa - Mahal Listesi.xlsx"))
    print("kaydedildi:", p)
